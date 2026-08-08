import logging
import os

import backoff
import pandas as pd
from google.api_core.exceptions import ResourceExhausted
from langchain_core.language_models.chat_models import BaseChatModel
from langgraph.graph import END, StateGraph

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

from check_planner.agents.models import (
    LegislativeReference,
    RegulationControl,
    VerifiedAgentState,
)
from check_planner.agents.prompts import controle_systeme_prompt
from check_planner.llm import get_llm_gemini, get_llm_groq, llm_groq
from check_planner.retriever import retrieve_regulation


class VerifierAgent:

    def __init__(self, llm: BaseChatModel = llm_groq):
        self.llm_params = {"iteration": 0, "type": "groq", "max": 20, "call": "verify"}
        self.llm_gen = llm.with_structured_output(RegulationControl)
        self.controle_llm = llm.with_structured_output(LegislativeReference)
        self.regulations = []

        # build graph
        self.graph = self._build_graph()

        # display(Image(self.graph.get_graph().draw_mermaid_png()))

    def _build_graph(self):
        """Build the LangGraph state graph"""
        graph = StateGraph(VerifiedAgentState)

        # add nodes
        graph.add_node("start", self._start_node)
        graph.add_node("load_check", self._load_check_plan_node)
        graph.add_node("plan", self._process_plan_node)
        graph.add_node("correction", self._line_correction_node)
        graph.add_node("stock", self._stock_data_node)
        graph.add_node("finish", self._finish_node)

        # add edges
        graph.add_edge("start", "load_check")
        graph.add_edge("load_check", "plan")

        graph.add_conditional_edges(
            "plan",
            self._should_regulation_continu,
            {"correction": "correction", "plan": "plan", "stock": "stock"},
        )

        graph.add_edge("correction", "plan")
        graph.add_edge("stock", "finish")

        graph.add_edge("finish", END)
        # set entry point
        graph.set_entry_point("start")

        return graph.compile()

    def _start_node(self, state: VerifiedAgentState):

        return {**state, "current_rg_num": 0, "regulation": {}}

    async def _load_check_plan_node(self, state: VerifiedAgentState):
        logger.info("Chargement du plan de vérification...")
        check_file = state["check_path"]

        if os.path.exists(check_file):

            df_data = pd.read_excel(check_file)

            data = df_data.fillna("").to_dict(orient="records")

            self.regulations = data

            return {**state, "current_rg_num": 0, "max_regs": len(data)}

        return {**state}

    async def _process_plan_node(self, state: VerifiedAgentState):
        logger.info("Plan de process")

        return {**state, "current_rg_num": state.get("current_rg_num") + 1}

    async def _line_correction_node(self, state: VerifiedAgentState):
        logger.info("correction...")

        reg = self.regulations[state["current_rg_num"] - 1]

        if not all(reg.values()):

            self.regulations[state["current_rg_num"] - 1] = {}

            return state

        controle = f"""
        Article / Objet du Contrôle: {reg['Article / Objet du Contrôle']}
        Objectif: {reg['Objectif']}
        Documents de Référence: {reg['Documents de Référence']}
        Fréquence: {reg['Fréquence']}
        Critères de Conformité: {reg['Critères de Conformité']}
        Documents Requis: {reg['Documents Requis']}
        Détails et Explications pour le Contrôleur: {reg['Détails et Explications pour le Contrôleur']}
        Points spécifiques à contrôler avec détails: {reg['Points spécifiques à contrôler avec détails']}
        """

        try:
            retrieved = await retrieve_regulation(controle)

            self.llm_params["call"] = "verify"

            LReference = await self._safe_invoke(
                controle_systeme_prompt.format(controle=controle, retrieved=retrieved)
            )
            if not isinstance(LReference, LegislativeReference):
                raise ValueError("Reference recuperée n'est pas la forme entendu")

            self.regulations[state["current_rg_num"] - 1][
                "Reference de Legislative (AMMC)"
            ] = (
                LReference.reference_text
                if LReference.reference_text
                else "Aucun passage réglementaire pertinent n'a été trouvé"
            )
            logger.info("reference ajoutée avec success")
            return state

        except Exception as e:
            logger.info(f"Erreur correction node: {e}")
            return state

    def _should_regulation_continu(self, state: VerifiedAgentState):

        if state["current_rg_num"] >= state["max_regs"]:
            return "stock"

        return "correction"

    async def _stock_data_node(self, state: VerifiedAgentState):
        logger.info("stockage...")
        output = state["output_file"]

        if self.regulations:
            try:

                data = [
                    {**line, "N° de Contrôle": idx + 1}
                    for idx, line in enumerate(self.regulations)
                    if line
                ]
                df = pd.DataFrame(data)

                # Exporter en Excel
                df.to_excel(output, index=False)

                return state
            except Exception as e:
                logger.error(f"Erreur au niveau de stockage: {e}")
        return state

    async def _excel_formatter(self, file_path):
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        # Chargement
        wb = load_workbook(file_path)
        ws = wb.active

        # Renommer la feuille
        ws.title = "Plan de controle"

        # Ajuster largeur de colonnes
        col_widths = {
            "A": 12,
            "B": 33,
            "C": 42,
            "D": 32,
            "E": 21,
            "F": 40,
            "G": 37,
            "H": 63,
            "I": 69,
            "J": 69,
        }
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        # Ajuster hauteur des lignes
        # head

        ws.row_dimensions[1].height = 30

        # other lines
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 89

        # Styles
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(
            start_color="628e3d", end_color="628e3d", fill_type="solid"
        )
        alternate_fill = PatternFill(
            start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
        )  # gris clair élégant
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Appliquer styles
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            for cell in row:
                # Bordures et wrap text
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                cell.border = border

                if row_idx == 1:  # En-tête
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                else:  # Corps du tableau
                    if row_idx % 2 == 0:  # lignes paires uniquement
                        cell.fill = alternate_fill

        # Ajouter filtre automatique
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"

        # Sauvegarder
        wb.save(file_path)

    async def _finish_node(self, state: VerifiedAgentState):
        logger.info("Fin du process...")
        output = state["output_file"]

        if os.path.exists(output):
            await self._excel_formatter(output)

        return state

    async def arun(self, check_path: str):
        """
        Agent start flow
        """
        folder = os.path.dirname(check_path)  # Chemin du dossier
        filename = os.path.basename(check_path)
        if "plan_de_controle_" in filename:
            filename = filename.replace("plan_de_controle_", "plan_de_controle_final_")
        else:
            filename = f"plan_de_controle_final_{filename}"
        init_state = {
            "check_path": check_path,
            "current_rg_num": 0,
            "regulation": {},
            "max_regs": 0,
            "output_file": f"{folder }/{filename.lower()}",
        }
        try:
            return await self.graph.ainvoke(init_state, {"recursion_limit": 10000000})
        except Exception as e:
            logger.error(f"Erreur lors de l'execution de l'agent: {e}")
            return {"output_file": f"{folder }/{filename.lower()}", "error": str(e)}

    async def _rotate_llm(self):
        if self.llm_params["iteration"] >= self.llm_params["max"]:
            self.llm_params["iteration"] = 0

            self.llm_params["type"] = (
                "gemini" if self.llm_params["type"] == "groq" else "groq"
            )

        if self.llm_params["type"] == "groq":
            llm = get_llm_groq()
        else:
            llm = get_llm_gemini()

        self.llm_gen = llm.with_structured_output(RegulationControl)
        self.controle_llm = llm.with_structured_output(LegislativeReference)

    @backoff.on_exception(
        backoff.expo, (ResourceExhausted, Exception), max_tries=7, jitter=None
    )
    async def _safe_invoke(self, full_messages):
        try:
            self.llm_params["iteration"] += 1
            if self.llm_params.get("call") == "verify":
                return await self.controle_llm.ainvoke(full_messages)
            else:
                return await self.llm_gen.ainvoke(full_messages)
        except ResourceExhausted as e:
            logger.warning("[Quota] Clé API dépassée, on change...")
            await self._rotate_llm()
            raise e
        except Exception as e:
            logger.error(f"[Erreur] {e}, on essaye un autre LLM...")
            await self._rotate_llm()
            raise e
